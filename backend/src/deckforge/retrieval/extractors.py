"""Text extraction from uploaded files.

Extractors are registered by file extension, so supporting a new document type
is one class plus one registration — no changes to the upload service.
"""

from __future__ import annotations

import abc
import csv
import io
from dataclasses import dataclass, field
from pathlib import Path

from deckforge.core.errors import UnsupportedFileError
from deckforge.core.logging import get_logger
from deckforge.core.registry import Registry

log = get_logger(__name__)


@dataclass(slots=True)
class ExtractedPage:
    """One logical page/sheet/slide of extracted text."""

    text: str
    page: int | None = None
    label: str = ""


@dataclass(slots=True)
class Extraction:
    """The full result of reading a document."""

    pages: list[ExtractedPage] = field(default_factory=list)
    kind: str = "document"
    metadata: dict[str, str] = field(default_factory=dict)

    @property
    def text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text.strip())

    def excerpt(self, limit: int = 600) -> str:
        body = self.text.strip().replace("\r", "")
        return body[:limit] + ("…" if len(body) > limit else "")


class Extractor(abc.ABC):
    """Reads one family of file types."""

    #: Lowercase extensions, without the dot.
    extensions: tuple[str, ...] = ()
    kind: str = "document"

    @abc.abstractmethod
    def extract(self, path: Path) -> Extraction:
        """Read ``path`` and return its text."""


EXTRACTORS: Registry[Extractor] = Registry("extractor")


def register_extractor(extractor: Extractor, *, override: bool = True) -> Extractor:
    """Register ``extractor`` for each of its extensions."""
    for extension in extractor.extensions:
        EXTRACTORS.register(extension, extractor, override=override)
    return extractor


def extract_file(path: Path) -> Extraction:
    """Extract text from ``path``.

    Raises:
        UnsupportedFileError: when no extractor handles the extension.
    """
    extension = path.suffix.lower().lstrip(".")
    extractor = EXTRACTORS.try_get(extension)
    if extractor is None:
        raise UnsupportedFileError(
            f"no extractor for '.{extension}' files",
            details={"supported": EXTRACTORS.names()},
        )
    try:
        return extractor.extract(path)
    except UnsupportedFileError:
        raise
    except Exception as exc:
        log.warning("extract.failed", path=path.name, error=str(exc))
        raise UnsupportedFileError(f"could not read {path.name}: {exc}") from exc


# --------------------------------------------------------------------------- #
# Built-in extractors
# --------------------------------------------------------------------------- #


class PlainTextExtractor(Extractor):
    extensions = ("txt", "md", "markdown", "rst", "log", "json", "yaml", "yml")
    kind = "document"

    def extract(self, path: Path) -> Extraction:
        text = path.read_text(encoding="utf-8", errors="replace")
        return Extraction(pages=[ExtractedPage(text=text)], kind=self.kind)


class PdfExtractor(Extractor):
    extensions = ("pdf",)
    kind = "document"

    def extract(self, path: Path) -> Extraction:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = [
            ExtractedPage(text=(page.extract_text() or "").strip(), page=i + 1)
            for i, page in enumerate(reader.pages)
        ]
        metadata = {}
        if reader.metadata:
            metadata = {
                k.lstrip("/"): str(v) for k, v in reader.metadata.items() if isinstance(v, str)
            }
        return Extraction(pages=pages, kind=self.kind, metadata=metadata)


class DocxExtractor(Extractor):
    extensions = ("docx",)
    kind = "document"

    def extract(self, path: Path) -> Extraction:
        from docx import Document

        document = Document(str(path))
        blocks = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
        return Extraction(pages=[ExtractedPage(text="\n".join(blocks))], kind=self.kind)


class PptxExtractor(Extractor):
    extensions = ("pptx",)
    kind = "presentation"

    def extract(self, path: Path) -> Extraction:
        from pptx import Presentation as PptxPresentation

        deck = PptxPresentation(str(path))
        pages: list[ExtractedPage] = []
        for i, slide in enumerate(deck.slides):
            lines: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    lines.append(shape.text_frame.text.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        lines.append(" | ".join(c.text.strip() for c in row.cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                lines.append(f"[notes] {slide.notes_slide.notes_text_frame.text.strip()}")
            pages.append(ExtractedPage(text="\n".join(lines), page=i + 1, label=f"slide {i + 1}"))
        return Extraction(pages=pages, kind=self.kind)


class CsvExtractor(Extractor):
    extensions = ("csv", "tsv")
    kind = "data"

    def extract(self, path: Path) -> Extraction:
        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        raw = path.read_text(encoding="utf-8", errors="replace")
        reader = csv.reader(io.StringIO(raw), delimiter=delimiter)
        rows = [" | ".join(row) for row in reader]
        return Extraction(pages=[ExtractedPage(text="\n".join(rows))], kind=self.kind)


class ExcelExtractor(Extractor):
    extensions = ("xlsx", "xlsm")
    kind = "data"

    def extract(self, path: Path) -> Extraction:
        from openpyxl import load_workbook

        workbook = load_workbook(str(path), read_only=True, data_only=True)
        pages: list[ExtractedPage] = []
        for index, sheet in enumerate(workbook.worksheets):
            rows = [
                " | ".join("" if c is None else str(c) for c in row)
                for row in sheet.iter_rows(values_only=True)
            ]
            pages.append(
                ExtractedPage(text="\n".join(rows), page=index + 1, label=str(sheet.title))
            )
        workbook.close()
        return Extraction(pages=pages, kind=self.kind)


class ImageExtractor(Extractor):
    """Images carry no text, but recording their dimensions makes them usable as assets."""

    extensions = ("png", "jpg", "jpeg", "gif", "webp", "bmp", "svg")
    kind = "image"

    def extract(self, path: Path) -> Extraction:
        metadata: dict[str, str] = {}
        if path.suffix.lower() != ".svg":
            try:
                from PIL import Image

                with Image.open(path) as image:
                    metadata = {
                        "width": str(image.width),
                        "height": str(image.height),
                        "format": image.format or "",
                    }
            except Exception as exc:  # pragma: no cover - depends on file
                log.debug("extract.image_metadata_failed", error=str(exc))
        return Extraction(
            pages=[ExtractedPage(text=f"[image] {path.name}")], kind=self.kind, metadata=metadata
        )


for _extractor in (
    PlainTextExtractor(),
    PdfExtractor(),
    DocxExtractor(),
    PptxExtractor(),
    CsvExtractor(),
    ExcelExtractor(),
    ImageExtractor(),
):
    register_extractor(_extractor)


def supported_extensions() -> list[str]:
    return EXTRACTORS.names()
