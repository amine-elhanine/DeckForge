"""PowerPoint (.pptx) exporter.

Generated entirely with ``python-pptx`` — no PowerPoint, no LibreOffice, no
network. Geometry comes from the same layout frames the HTML preview uses, so
the exported file matches what the user approved on screen.

Coordinates: one logical canvas unit is one CSS pixel is 9525 EMU, which makes a
1280×720 deck exactly 13.333in × 7.5in.
"""

from __future__ import annotations

import io
from pathlib import Path

from pptx import Presentation as PptxPresentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.shapes.autoshape import Shape as AutoShape
from pptx.slide import Slide as PptxSlide
from pptx.util import Emu, Pt

from deckforge.core.logging import get_logger
from deckforge.exporters.base import ExportContext, Exporter
from deckforge.layouts.model import DecorPlacement, ElementPlacement, LayoutFrame, TextPlacement
from deckforge.models.deck import (
    Background,
    Box,
    BulletsElement,
    CardsElement,
    ChartElement,
    CodeElement,
    DiagramElement,
    IconElement,
    ImageElement,
    MetricElement,
    Presentation,
    QuoteElement,
    ShapeElement,
    TableElement,
    TextElement,
    TimelineElement,
)
from deckforge.models.enums import Align, BackgroundKind, BulletStyle, ChartKind
from deckforge.renderers import diagrams, icons
from deckforge.renderers.charts import series_colors
from deckforge.renderers.html import MARKER_GLYPHS
from deckforge.renderers.inline import parse_inline, plain
from deckforge.renderers.paths import flatten_path, scale_subpaths
from deckforge.renderers.primitives import (
    Drawing,
    Ellipse,
    Label,
    Line,
    Polyline,
    Rect,
    Wedge,
)
from deckforge.themes.model import Theme, TypeStyle
from deckforge.utils import colors

log = get_logger(__name__)

EMU_PER_UNIT = 9525
"""EMU per logical canvas unit (= per CSS pixel at 96 dpi)."""
PT_PER_UNIT = 0.75
"""Points per logical unit."""

ALIGNMENT = {
    Align.START: PP_ALIGN.LEFT,
    Align.CENTER: PP_ALIGN.CENTER,
    Align.END: PP_ALIGN.RIGHT,
    Align.JUSTIFY: PP_ALIGN.JUSTIFY,
}

CHART_TYPES = {
    ChartKind.COLUMN: XL_CHART_TYPE.COLUMN_CLUSTERED,
    ChartKind.BAR: XL_CHART_TYPE.BAR_CLUSTERED,
    ChartKind.STACKED_BAR: XL_CHART_TYPE.COLUMN_STACKED,
    ChartKind.LINE: XL_CHART_TYPE.LINE_MARKERS,
    ChartKind.AREA: XL_CHART_TYPE.AREA,
    ChartKind.PIE: XL_CHART_TYPE.PIE,
    ChartKind.DONUT: XL_CHART_TYPE.DOUGHNUT,
    ChartKind.SCATTER: XL_CHART_TYPE.XY_SCATTER,
    ChartKind.RADAR: XL_CHART_TYPE.RADAR_MARKERS,
}


def emu(value: float) -> Emu:
    return Emu(round(value * EMU_PER_UNIT))


def pt(value: float) -> Pt:
    return Pt(max(1.0, value * PT_PER_UNIT))


def rgb(value: str) -> RGBColor:
    """Convert a hex or ``rgba()`` colour to a PowerPoint RGB value."""
    try:
        r, g, b = colors.to_rgb_tuple(value)
    except ValueError:
        r, g, b = (17, 19, 24)
    return RGBColor(r, g, b)


class PptxExporter(Exporter):
    """Renders a deck as an editable PowerPoint file."""

    name = "pptx"
    label = "PowerPoint"
    extension = "pptx"
    media_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    binary = True
    description = "Fully editable .pptx with native charts, tables and speaker notes."

    def render(self, deck: Presentation, context: ExportContext) -> bytes:
        theme = context.theme
        canvas_w, canvas_h = deck.canvas
        prs = PptxPresentation()
        prs.slide_width = emu(canvas_w)
        prs.slide_height = emu(canvas_h)
        blank = prs.slide_layouts[6]

        frames = context.renderer.layouts.resolve_deck(deck, theme)
        for slide_model, frame in zip(deck.slides, frames, strict=True):
            if slide_model.hidden and not context.option("include_hidden", False):
                continue
            pptx_slide = prs.slides.add_slide(blank)
            self._background(pptx_slide, frame.background, theme)
            for decoration in frame.decor:
                self._decor(pptx_slide, decoration)
            # Backdrops (z < 0) go down first so titles stay legible on top of them.
            for placement in (p for p in frame.elements if p.z < 0):
                self._element(pptx_slide, placement, theme, frame, context)
            for text in frame.texts:
                self._text_placement(pptx_slide, text, theme, frame)
            for placement in (p for p in frame.elements if p.z >= 0):
                self._element(pptx_slide, placement, theme, frame, context)
            if slide_model.notes:
                pptx_slide.notes_slide.notes_text_frame.text = slide_model.notes
            if slide_model.references:
                notes = pptx_slide.notes_slide.notes_text_frame
                notes.add_paragraph().text = "References: " + "; ".join(
                    r.format_apa() for r in slide_model.references
                )

        buffer = io.BytesIO()
        prs.save(buffer)
        return buffer.getvalue()

    # -- background & decoration -------------------------------------------- #

    def _background(self, slide: PptxSlide, background: Background, theme: Theme) -> None:
        fill = slide.background.fill
        match background.kind:
            case BackgroundKind.GRADIENT | BackgroundKind.MESH:
                stops = background.colors or [theme.palette.background, theme.palette.surface]
                fill.gradient()
                fill.gradient_angle = background.angle
                gradient_stops = fill.gradient_stops
                gradient_stops[0].color.rgb = rgb(stops[0])
                gradient_stops[1].color.rgb = rgb(stops[-1])
            case BackgroundKind.NONE:
                return
            case _:
                fill.solid()
                fill.fore_color.rgb = rgb(background.color or theme.palette.background)

    def _decor(self, slide: PptxSlide, decoration: DecorPlacement) -> None:
        shapes = {
            "rect": MSO_SHAPE.RECTANGLE,
            "rule": MSO_SHAPE.RECTANGLE,
            "ellipse": MSO_SHAPE.OVAL,
            "triangle": MSO_SHAPE.RIGHT_TRIANGLE,
            "diagonal": MSO_SHAPE.RIGHT_TRIANGLE,
        }
        shape_type = shapes.get(decoration.kind)
        if shape_type is None:  # grids and dot fields are decorative noise in PPTX
            return
        b = decoration.box
        shape = slide.shapes.add_shape(
            shape_type, emu(max(0, b.x)), emu(max(0, b.y)), emu(b.width), emu(b.height)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(decoration.color)
        shape.fill.transparency = 1.0 - decoration.opacity
        shape.line.fill.background()
        shape.shadow.inherit = False

    # -- text ---------------------------------------------------------------- #

    def _textbox(self, slide: PptxSlide, box: Box, *, anchor: MSO_ANCHOR = MSO_ANCHOR.TOP):
        shape = slide.shapes.add_textbox(emu(box.x), emu(box.y), emu(box.width), emu(box.height))
        frame = shape.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = anchor
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        return shape, frame

    def _apply_style(
        self, paragraph, style: TypeStyle, theme: Theme, *, invert: bool, align: Align | None = None
    ) -> None:
        paragraph.line_spacing = style.line_height
        if align is not None:
            paragraph.alignment = ALIGNMENT[align]
        colour = (
            "#ffffff" if invert and style.color in ("text", "heading") else theme.color(style.color)
        )
        for run in paragraph.runs:
            run.font.size = pt(style.size)
            run.font.bold = run.font.bold or style.weight >= 600
            run.font.name = theme.office_font(style.font)
            run.font.color.rgb = rgb(colour)

    def _write(
        self,
        frame,
        text: str,
        style: TypeStyle,
        theme: Theme,
        *,
        invert: bool,
        align: Align = Align.START,
        first: bool = True,
    ) -> None:
        """Write ``text`` into ``frame`` as styled runs, one paragraph per line."""
        for i, line in enumerate(text.split("\n")):
            paragraph = frame.paragraphs[0] if (first and i == 0) else frame.add_paragraph()
            for span in parse_inline(line) or [
                type("S", (), {"text": "", "bold": False, "italic": False, "code": False})()
            ]:
                run = paragraph.add_run()
                run.text = span.text.upper() if style.transform == "uppercase" else span.text
                run.font.bold = bool(span.bold)
                run.font.italic = bool(span.italic)
                if getattr(span, "code", False):
                    run.font.name = theme.fonts.mono
            self._apply_style(paragraph, style, theme, invert=invert, align=align)

    def _text_placement(
        self, slide: PptxSlide, placement: TextPlacement, theme: Theme, frame: LayoutFrame
    ) -> None:
        anchor = {"start": MSO_ANCHOR.TOP, "center": MSO_ANCHOR.MIDDLE, "end": MSO_ANCHOR.BOTTOM}[
            placement.valign
        ]
        _, text_frame = self._textbox(slide, placement.box, anchor=anchor)
        style = theme.type_scale.get(placement.role)
        self._write(
            text_frame,
            placement.text,
            style,
            theme,
            invert=frame.invert_text,
            align=placement.align,
        )

    # -- elements ------------------------------------------------------------ #

    def _element(
        self,
        slide: PptxSlide,
        placement: ElementPlacement,
        theme: Theme,
        frame: LayoutFrame,
        context: ExportContext,
    ) -> None:
        element, box = placement.element, placement.box
        anchor = {"start": MSO_ANCHOR.TOP, "center": MSO_ANCHOR.MIDDLE, "end": MSO_ANCHOR.BOTTOM}[
            placement.valign
        ]
        invert = frame.invert_text
        # The layout may have shrunk the type to fit; every size below follows it.
        theme = theme.type_scaled(placement.scale)
        ts = theme.type_scale

        match element:
            case TextElement():
                _, tf = self._textbox(slide, box, anchor=anchor)
                self._write(
                    tf,
                    element.text,
                    ts.get(str(element.role)),
                    theme,
                    invert=invert,
                    align=element.align,
                )

            case BulletsElement():
                self._bullets(slide, element, box, theme, invert=invert, anchor=anchor)

            case ImageElement():
                self._image(slide, element, box, theme, context)

            case IconElement():
                self._icon(
                    slide,
                    element.name,
                    box.x,
                    box.y,
                    element.size,
                    theme.color(element.color or "primary"),
                    theme,
                )
                if element.label:
                    label_box = Box(
                        x=box.x,
                        y=box.y + element.size * 1.2,
                        width=box.width,
                        height=max(20.0, ts.caption.size * 2),
                    )
                    _, tf = self._textbox(slide, label_box)
                    self._write(tf, element.label, ts.caption, theme, invert=invert)

            case ChartElement():
                self._chart(slide, element, box, theme)

            case TableElement():
                self._table(slide, element, box, theme)

            case QuoteElement():
                _, tf = self._textbox(slide, box, anchor=MSO_ANCHOR.MIDDLE)
                self._write(tf, f"“{plain(element.text)}”", ts.quote, theme, invert=invert)
                if element.attribution:
                    role = f" · {element.role}" if element.role else ""
                    self._write(
                        tf,
                        f"— {element.attribution}{role}",
                        ts.caption,
                        theme,
                        invert=invert,
                        first=False,
                    )

            case MetricElement():
                _, tf = self._textbox(slide, box, anchor=MSO_ANCHOR.MIDDLE)
                self._write(tf, element.value, ts.metric_value, theme, invert=False)
                self._write(tf, element.label, ts.metric_label, theme, invert=invert, first=False)
                if element.delta:
                    delta_style = ts.metric_label.model_copy(
                        update={
                            "color": {"up": "success", "down": "danger"}.get(
                                element.trend or "", "text_muted"
                            )
                        }
                    )
                    self._write(tf, element.delta, delta_style, theme, invert=False, first=False)

            case CardsElement():
                self._cards(slide, element, box, theme)

            case TimelineElement():
                self._timeline(slide, element, box, theme, invert=invert)

            case CodeElement():
                shape = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE,
                    emu(box.x),
                    emu(box.y),
                    emu(box.width),
                    emu(box.height),
                )
                shape.fill.solid()
                shape.fill.fore_color.rgb = rgb(theme.palette.surface)
                shape.line.color.rgb = rgb(theme.palette.border)
                shape.shadow.inherit = False
                inner = box.inset(16, 14)
                _, tf = self._textbox(slide, inner)
                self._write(tf, element.source, ts.code, theme, invert=False)

            case DiagramElement():
                # Native shapes, not a picture: the diagram stays editable in
                # PowerPoint and stays sharp at any zoom.
                drawing = diagrams.render_diagram(element, box.width, box.height, theme)
                if drawing is not None:
                    self._draw(slide, drawing.translated(box.x, box.y), theme)
                    return

                rendered = (
                    context.resolve_asset(element.rendered_asset_id)
                    if element.rendered_asset_id
                    else None
                )
                if rendered is not None:
                    self._picture(slide, rendered, box)
                else:
                    _, tf = self._textbox(slide, box)
                    self._write(tf, element.source, ts.code, theme, invert=False)

            case ShapeElement():
                shape_type = MSO_SHAPE.OVAL if element.shape == "ellipse" else MSO_SHAPE.RECTANGLE
                shape = slide.shapes.add_shape(
                    shape_type, emu(box.x), emu(box.y), emu(box.width), emu(box.height)
                )
                shape.fill.solid()
                shape.fill.fore_color.rgb = rgb(theme.color(element.color or "primary"))
                shape.fill.transparency = 1.0 - element.opacity
                shape.line.fill.background()
                shape.shadow.inherit = False

            case _:
                return

    def _bullets(
        self,
        slide: PptxSlide,
        element: BulletsElement,
        box: Box,
        theme: Theme,
        *,
        invert: bool,
        anchor: MSO_ANCHOR,
    ) -> None:
        _, tf = self._textbox(slide, box, anchor=anchor)
        style = theme.type_scale.bullet
        glyph = MARKER_GLYPHS.get(theme.bullets.marker, "•")
        for i, item in enumerate(element.items):
            paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            paragraph.level = min(4, item.level)
            paragraph.space_after = pt(theme.bullets.row_gap)
            marker = (
                f"{i + 1}. "
                if element.bullet_style is BulletStyle.NUMBER
                else ("" if element.bullet_style is BulletStyle.NONE else f"{glyph}  ")
            )
            if marker:
                run = paragraph.add_run()
                run.text = marker
                run.font.color.rgb = rgb(theme.color(theme.bullets.marker_color))
            for span in parse_inline(item.text):
                run = paragraph.add_run()
                run.text = span.text
                run.font.bold = span.bold or item.emphasis
                run.font.italic = span.italic
            self._apply_style(paragraph, style, theme, invert=invert)

    def _image(
        self,
        slide: PptxSlide,
        element: ImageElement,
        box: Box,
        theme: Theme,
        context: ExportContext,
    ) -> None:
        path = context.resolve_asset(element.src)
        if path is not None:
            self._picture(slide, path, box)
            return
        # No local bytes: draw a labelled placeholder rather than failing the export.
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, emu(box.x), emu(box.y), emu(box.width), emu(box.height)
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(theme.palette.surface)
        shape.line.color.rgb = rgb(theme.palette.border)
        shape.line.dash_style = None
        shape.shadow.inherit = False
        tf = shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = tf.paragraphs[0]
        run = paragraph.add_run()
        run.text = element.alt or element.caption or "Image"
        paragraph.alignment = PP_ALIGN.CENTER
        self._apply_style(paragraph, theme.type_scale.caption, theme, invert=False)

    @staticmethod
    def _picture(slide: PptxSlide, path: Path, box: Box) -> None:
        try:
            slide.shapes.add_picture(
                str(path), emu(box.x), emu(box.y), emu(box.width), emu(box.height)
            )
        except Exception as exc:  # pragma: no cover - depends on file contents
            log.warning("pptx.image_failed", path=str(path), error=str(exc))

    def _icon(
        self,
        slide: PptxSlide,
        name: str | None,
        x: float,
        y: float,
        size: float,
        colour: str,
        theme: Theme,
    ) -> None:
        """Stroke an icon as freeform shapes flattened from its SVG path."""
        try:
            subpaths = scale_subpaths(flatten_path(icons.get_icon(name)), size, dx=x, dy=y)
        except Exception:  # pragma: no cover - malformed plugin icon
            return
        for points in subpaths:
            if len(points) < 2:
                continue
            builder = slide.shapes.build_freeform(emu(points[0][0]), emu(points[0][1]), scale=1.0)
            builder.add_line_segments([(emu(px), emu(py)) for px, py in points[1:]], close=False)
            shape = builder.convert_to_shape()
            shape.fill.background()
            shape.line.color.rgb = rgb(colour)
            shape.line.width = pt(theme.icon_stroke)
            shape.shadow.inherit = False

    def _draw(self, slide: PptxSlide, drawing: Drawing, theme: Theme) -> None:
        """Replay vector primitives as native PowerPoint shapes.

        Charts go in as real Office charts, but a flowchart has no native
        equivalent — drawing it out of shapes keeps it editable and vector-sharp
        instead of pasting in a bitmap.
        """
        for item in drawing.items:
            match item:
                case Rect():
                    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if item.radius else MSO_SHAPE.RECTANGLE
                    shape = slide.shapes.add_shape(
                        shape_type,
                        emu(item.x),
                        emu(item.y),
                        emu(max(1.0, item.width)),
                        emu(max(1.0, item.height)),
                    )
                    self._style_shape(
                        shape, item.fill, item.stroke, item.stroke_width, item.opacity
                    )
                case Ellipse():
                    shape = slide.shapes.add_shape(
                        MSO_SHAPE.OVAL,
                        emu(item.cx - item.rx),
                        emu(item.cy - item.ry),
                        emu(max(1.0, item.rx * 2)),
                        emu(max(1.0, item.ry * 2)),
                    )
                    self._style_shape(
                        shape, item.fill, item.stroke, item.stroke_width, item.opacity
                    )
                case Line():
                    connector = slide.shapes.add_connector(
                        MSO_CONNECTOR.STRAIGHT,
                        emu(item.x1),
                        emu(item.y1),
                        emu(item.x2),
                        emu(item.y2),
                    )
                    connector.line.color.rgb = rgb(item.stroke)
                    connector.line.width = pt(max(0.5, item.stroke_width))
                case Polyline():
                    self._freeform(slide, item)
                case Wedge():
                    self._freeform(
                        slide,
                        Polyline(
                            item.outline(), fill=item.fill, closed=True, stroke=None, stroke_width=0
                        ),
                    )
                case Label():
                    self._draw_label(slide, item, theme)

    @staticmethod
    def _style_shape(
        shape: AutoShape, fill: str | None, stroke: str | None, width: float, opacity: float
    ) -> None:
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(fill)
            if opacity < 1:
                shape.fill.transparency = 1.0 - opacity
        else:
            shape.fill.background()
        if stroke and width:
            shape.line.color.rgb = rgb(stroke)
            shape.line.width = pt(max(0.5, width))
        else:
            shape.line.fill.background()
        shape.shadow.inherit = False

    def _freeform(self, slide: PptxSlide, item: Polyline) -> None:
        points = item.points
        if len(points) < 2:
            return
        builder = slide.shapes.build_freeform(emu(points[0][0]), emu(points[0][1]), scale=1.0)
        builder.add_line_segments([(emu(x), emu(y)) for x, y in points[1:]], close=item.closed)
        shape = builder.convert_to_shape()
        self._style_shape(shape, item.fill, item.stroke, item.stroke_width, item.opacity)

    def _draw_label(self, slide: PptxSlide, item: Label, theme: Theme) -> None:
        # Labels arrive positioned by their anchor point, so the box is centred
        # on it and the paragraph alignment does the rest.
        width = max(40.0, len(item.text) * item.size * 0.72 + 16)
        height = item.size * 2.0
        left = {"start": item.x, "middle": item.x - width / 2, "end": item.x - width}[item.anchor]
        top = {"top": item.y, "middle": item.y - height / 2, "bottom": item.y - height}[
            item.baseline
        ]
        _, frame = self._textbox(
            slide, Box(x=left, y=top, width=width, height=height), anchor=MSO_ANCHOR.MIDDLE
        )
        style = theme.type_scale.caption.model_copy(
            update={
                "size": item.size,
                "weight": item.weight,
                "font": item.family,
                "color": item.color,
            }
        )
        self._write(frame, item.text, style, theme, invert=False, align=Align.CENTER)

    def _chart(self, slide: PptxSlide, element: ChartElement, box: Box, theme: Theme) -> None:
        spec = element.chart
        if not spec.series:
            return
        data = CategoryChartData()
        data.categories = spec.categories or [
            f"#{i + 1}" for i in range(len(spec.series[0].values))
        ]
        for series in spec.series:
            data.add_series(series.name, tuple(series.values))
        chart_type = CHART_TYPES.get(spec.kind, XL_CHART_TYPE.COLUMN_CLUSTERED)
        graphic = slide.shapes.add_chart(
            chart_type, emu(box.x), emu(box.y), emu(box.width), emu(box.height), data
        )
        chart = graphic.chart
        chart.has_title = bool(element.title)
        if element.title:
            chart.chart_title.text_frame.text = element.title
        chart.has_legend = spec.legend and len(spec.series) > 1
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        palette = series_colors(spec, theme)
        for i, plot_series in enumerate(chart.series):
            try:
                plot_series.format.fill.solid()
                plot_series.format.fill.fore_color.rgb = rgb(palette[i % len(palette)])
            except AttributeError, ValueError:  # line/scatter types style differently
                continue

    def _table(self, slide: PptxSlide, element: TableElement, box: Box, theme: Theme) -> None:
        columns = element.columns or (element.rows[0] if element.rows else [])
        if not columns:
            return
        row_count = len(element.rows) + (1 if element.header else 0)
        graphic = slide.shapes.add_table(
            max(1, row_count), len(columns), emu(box.x), emu(box.y), emu(box.width), emu(box.height)
        )
        table = graphic.table
        style = theme.type_scale.table
        offset = 0
        if element.header:
            offset = 1
            for c, label in enumerate(columns):
                cell = table.cell(0, c)
                cell.text = plain(label)
                for paragraph in cell.text_frame.paragraphs:
                    self._apply_style(paragraph, style, theme, invert=True)
                    for run in paragraph.runs:
                        run.font.bold = True
        for r, row in enumerate(element.rows):
            for c in range(len(columns)):
                cell = table.cell(r + offset, c)
                cell.text = plain(row[c]) if c < len(row) else ""
                for paragraph in cell.text_frame.paragraphs:
                    self._apply_style(paragraph, style, theme, invert=False)

    def _cards(self, slide: PptxSlide, element: CardsElement, box: Box, theme: Theme) -> None:
        from deckforge.layouts.builtin import grid_cells

        cells = grid_cells(box, len(element.cards), max(1, element.columns), theme.spacing.gap)
        for card, cell in zip(element.cards, cells, strict=True):
            accent = theme.color(card.accent or "primary")
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                emu(cell.x),
                emu(cell.y),
                emu(cell.width),
                emu(cell.height),
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = rgb(theme.color(theme.shape.card_fill))
            shape.line.color.rgb = rgb(theme.palette.border)
            shape.shadow.inherit = False
            slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, emu(cell.x), emu(cell.y), emu(cell.width), emu(3)
            ).fill.solid()
            accent_bar = slide.shapes[-1]
            accent_bar.fill.fore_color.rgb = rgb(accent)
            accent_bar.line.fill.background()
            accent_bar.shadow.inherit = False

            inner = cell.inset(theme.spacing.card_padding, theme.spacing.card_padding * 0.8)
            if card.icon:
                self._icon(slide, card.icon, inner.x, inner.y, 26, accent, theme)
                inner = Box(
                    x=inner.x,
                    y=inner.y + 38,
                    width=inner.width,
                    height=max(30.0, inner.height - 38),
                )
            # Centred, to match the HTML preview: a card taller than its copy
            # looks deliberate with the text in the middle and unfinished with
            # it pinned to the top.
            _, tf = self._textbox(slide, inner, anchor=MSO_ANCHOR.MIDDLE)
            self._write(tf, card.title, theme.type_scale.card_title, theme, invert=False)
            if card.body:
                self._write(
                    tf, card.body, theme.type_scale.card_body, theme, invert=False, first=False
                )

    def _timeline(
        self, slide: PptxSlide, element: TimelineElement, box: Box, theme: Theme, *, invert: bool
    ) -> None:
        from deckforge.layouts.builtin import columns as split_columns
        from deckforge.layouts.builtin import rows as split_rows

        entries = element.entries
        if not entries:
            return
        horizontal = element.orientation == "horizontal"
        cells = (
            split_columns(box, len(entries), theme.spacing.gap)
            if horizontal
            else split_rows(box, len(entries), theme.spacing.tight_gap)
        )
        rail_colour = rgb(theme.palette.border)
        if horizontal:
            rail = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, emu(box.x), emu(box.y + 7), emu(box.width), emu(2)
            )
        else:
            rail = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, emu(box.x + 6), emu(box.y), emu(2), emu(box.height)
            )
        rail.fill.solid()
        rail.fill.fore_color.rgb = rail_colour
        rail.line.fill.background()
        rail.shadow.inherit = False

        for entry, cell in zip(entries, cells, strict=True):
            colour = theme.color(
                "success"
                if entry.status == "done"
                else ("text_muted" if entry.status == "planned" else "primary")
            )
            dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, emu(cell.x), emu(cell.y), emu(14), emu(14))
            dot.fill.solid()
            dot.fill.fore_color.rgb = rgb(colour)
            dot.line.fill.background()
            dot.shadow.inherit = False

            text_box = Box(
                x=cell.x,
                y=cell.y + 24,
                width=max(60.0, cell.width - theme.spacing.gap),
                height=max(40.0, cell.height - 24),
            )
            _, tf = self._textbox(slide, text_box)
            label_style = theme.type_scale.eyebrow.model_copy(update={"color": colour})
            self._write(tf, entry.label, label_style, theme, invert=False)
            self._write(
                tf, entry.title, theme.type_scale.card_title, theme, invert=invert, first=False
            )
            if entry.body:
                self._write(
                    tf, entry.body, theme.type_scale.card_body, theme, invert=invert, first=False
                )
